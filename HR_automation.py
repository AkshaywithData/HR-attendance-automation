import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl import load_workbook
import glob
import shutil
import os


def read_data():

    files = glob.glob("Data/Source Files/Attendance file/*.csv")
    files += glob.glob("Data/Source Files/Attendance file/*.xlsx")  

    if not files:
        raise FileNotFoundError("No CSV or Excel file found.")

    dataframes = []

    for file in files:
    
        if file.endswith(".csv"):
            df = pd.read_csv(file)  
        else:
            df = pd.read_excel(file)

        dataframes.append(df)

    df = pd.concat(dataframes, ignore_index=True)

    employee_master = pd.read_excel("Data/Source Files/Employees file/Employee_Master.xlsx")

    archive_folder = "Data/Archive"

    os.makedirs(archive_folder, exist_ok = True)

    for file in files:
        shutil.move(file, archive_folder)

    return df, employee_master

    
def clean_data(df,employee_master):

    df["Date"] = pd.to_datetime(df["Date"])

    df_long = df.melt(id_vars="Date",var_name="Employee",value_name="Time")

    df_long = df_long.merge(employee_master,on="Employee",how="left")

    df_long["Employee"] = df_long["Employee_Name"]

    df_long.drop(columns="Employee_Name", inplace=True)

    df_long[["In_Time", "Out_Time"]] = df_long["Time"].str.split("-", expand=True)

    df_long["In_Time"] = pd.to_datetime(df_long["In_Time"],format="%H:%M",errors="coerce")

    df_long["Out_Time"] = pd.to_datetime(df_long["Out_Time"],format="%H:%M",errors="coerce")

    df_long["Working_Hours"] = df_long["Out_Time"] - df_long["In_Time"]

    office_start = pd.to_timedelta("09:00:00")
    in_time = (
        pd.to_timedelta(df_long["In_Time"].dt.hour, unit="h") +                                                                 
        pd.to_timedelta(df_long["In_Time"].dt.minute, unit="m") +
        pd.to_timedelta(df_long["In_Time"].dt.second, unit="s")
    )
    df_long["Late_Minutes"] = ((in_time - office_start).dt.total_seconds().div(60).clip(lower=0))

    office_end   = pd.to_timedelta("17:00:00")
    out_time = (
        pd.to_timedelta(df_long["Out_Time"].dt.hour, unit="h") +
        pd.to_timedelta(df_long["Out_Time"].dt.minute, unit="m") +
    pd.to_timedelta(df_long["Out_Time"].dt.second, unit="s")
    )
    df_long["Early_Exit"] = ((office_end - out_time).dt.total_seconds().div(60).clip(lower=0))

    standard_hours = pd.to_timedelta("08:00:00")

    df_long["Overtime_Minutes"] = ((df_long["Working_Hours"] - standard_hours).dt.total_seconds().div(60).clip(lower=0))

    df_long["OT_Hours"] = df_long["Overtime_Minutes"].div(60).round(2)

    return df_long

def save_files(df_long):

    new_file = "Attendance_Combined_Cleaned.xlsx"

    os.makedirs("Data/Clean Files", exist_ok=True)

    df_long.to_excel(
        "Data/Clean Files/" + new_file,
        index=False
    )

def calculate_kpis(df_long):

    Total_employees = df_long["Employee"].nunique()

    df_long["Status"] = np.where(df_long["In_Time"].notna() & df_long["Out_Time"].notna(),
        "Present","Absent")

    total_days = df_long["Date"].nunique()

    total_present = (df_long["Status"] == "Present").sum()   

    total_absent = (df_long["Status"] == "Absent").sum()

    Attendance_percent = round((total_present / (total_present + total_absent)) * 100,2)

    avg_working_hours = df_long["Working_Hours"].mean().round("s")

    total_late_days = (df_long["Late_Minutes"] > 0).sum()

    total_ot_hours = round(df_long["Overtime_Minutes"].sum() / 60, 2)

    data = {
    "Total_employees": Total_employees,
    "total_present": total_present,
    "total_absent": total_absent,
    "Attendance_percent": Attendance_percent,
    "total_late_days": total_late_days,
    "total_ot_hours": total_ot_hours,
    "avg_working_hours": avg_working_hours,
    "total_days": total_days
    }
    return data

def get_issues(row):
    issues = []

    if row["Status"] == "Absent":
        issues.append("Absent")

    if row["Late_Minutes"] > 0:
        issues.append("late Arrival")

    if row["Working_Hours"] <  pd.Timedelta(hours = 8):
        issues.append("Short_Hours")

    if row["Overtime_Minutes"] > 0:
        issues.append("Overtime")

    if pd.isna(row["Out_Time"]) and pd.notna(row["In_Time"]):
               issues.append("Missing_punch_out")
    return ", ".join(issues)

def calculate_summaries_n_exception(df_long):

    employee_summary = (df_long.groupby("Employee").agg(
                        Present = ("Status", lambda x: (x == "Present").sum()),
                        Late = ("Late_Minutes", lambda x: (x >0).sum()),
                        OT = ("Overtime_Minutes", "sum"),
                        Avg_Hours = ("Working_Hours", "mean")
                        ).reset_index()
                        )
    employee_summary["OT"] = (employee_summary["OT"] / 60).round(2)

    employee_summary["Avg_Hours"] = (employee_summary["Avg_Hours"].dt.total_seconds() / 3600).round(2)

    daily_summary = (
        df_long.groupby("Date").agg(
            Present=("Status", lambda x: (x == "Present").sum()),
            Late=("Late_Minutes", lambda x: (x > 0).sum()),
            Avg_Hours=("Working_Hours", "mean")).reset_index()
    )
    daily_summary["Avg_Hours"] = (daily_summary["Avg_Hours"].dt.total_seconds() / 3600).round(2)

    overtime_summary = (
        df_long.groupby("Employee")
        .agg(
            OT_Hours=("OT_Hours", "sum"),
            OT_Days=("Overtime_Minutes", lambda x: (x > 0).sum())
        )
        .reset_index()
        .sort_values("OT_Hours", ascending=False)
    )

    df_long["Issue"] =  df_long.apply(get_issues, axis = 1)

    exception_report = df_long[df_long["Issue"] !=""]

    reports = {
    "employee": employee_summary,
    "daily": daily_summary,
    "exception": exception_report,
    "overtime": overtime_summary
    }
    return reports

def generate_graphs(df_long):

    df_long["Month"] = df_long["Date"].dt.to_period("M")

    monthly_ot = (df_long.groupby("Month")["Overtime_Minutes"].sum().reset_index())

    monthly_ot["OT_Hours"] = (monthly_ot["Overtime_Minutes"] / 60).round(2)

    top10_late = (
        df_long[df_long["Late_Minutes"] > 0]
        .groupby("Employee")
        .size()
        .reset_index(name="Late_Days")
        .sort_values("Late_Days", ascending=False)
        .head(10)
    )
    plt.figure(figsize=(10,5))
    plt.bar(top10_late["Employee"], top10_late["Late_Days"])
    plt.title("Top 10 Employees with Most Late Days")
    plt.xlabel("Employee")
    plt.ylabel("Late Days")
    plt.xticks(rotation=45)
    plt.tight_layout()
    plt.savefig("Generated Reports/Charts/Top 10 late employees.png", dpi = 200)
    plt.show()

    attendance_trend = (
        df_long.groupby("Date")["Status"]
        .apply(lambda x: (x == "Present").sum())
        .reset_index(name="Present")
    )
    top_ot_employees = (
        df_long.groupby("Employee")["OT_Hours"]
        .sum()
        .reset_index()
        .sort_values(by="OT_Hours", ascending=False)
        .head(10)
    )

    plt.figure(figsize=(8, 5))

    plt.barh(
        top_ot_employees["Employee"],
        top_ot_employees["OT_Hours"]
    )

    plt.title("Top 10 Employees by Overtime Hours")
    plt.xlabel("Overtime Hours")
    plt.ylabel("Employee")

    plt.gca().invert_yaxis()  # Highest OT at the top

    plt.tight_layout()
    plt.savefig("Generated Reports/Charts/top_ot_employees.png", dpi=200)
    plt.show()

    return monthly_ot, attendance_trend

def generate_reports_n_Dashboard(reports, data, attendance_trend, monthly_ot):
    wb = load_workbook("Generated Reports/HR_Attendance_Report.xlsx")

    # Delete old sheet
    if "Employee_Summary" in wb.sheetnames:
        del wb["Employee_Summary"]

    # Create sheet
    ws_emp = wb.create_sheet("Employee_Summary")

    # Report Title
    ws_emp["A1"] = "HR Attendance Report - Employee Summary"

    # Write dataframe
    for row in dataframe_to_rows(reports["employee"], index=False, header=True):
        ws_emp.append(row)

    if "Exception_Report" in wb.sheetnames:
        del wb["Exception_Report"]

    ws_exception = wb.create_sheet("Exception_Report")

    ws_exception["A1"] = "HR Attendance Report - Exception Report"

    for row in dataframe_to_rows(reports["exception"], index=False, header=True):
        ws_exception.append(row)

    if "Daily_Summary" in wb.sheetnames:
        del wb["Daily_Summary"]

    ws_daily = wb.create_sheet("Daily_Summary")

    ws_daily["A1"] = "HR Attendance Report - Daily Summary"

    for row in dataframe_to_rows(reports["daily"], index=False, header=True):
        ws_daily.append(row)

    if "Overtime_Summary" in wb.sheetnames:
        del wb["Overtime_Summary"]

    ws_ot = wb.create_sheet("Overtime_Summary")

    ws_ot["A1"] = "HR Attendance Report - Overtime Summary"

    for row in dataframe_to_rows(reports["overtime"], index=False, header=True):
        ws_ot.append(row)

    # Delete old Dashboard sheet if it exists
    if "Dashboard" in wb.sheetnames:
        del wb["Dashboard"]

    template = wb["Dashboard_Template"]

    ws = wb.copy_worksheet(template)
    ws.title = "Dashboard"

    # Create a new Dashboard sheet

    ws["B5"] = data["Total_employees"]
    ws["E5"] = data["total_present"]
    ws["H5"] = data["total_absent"]
    ws["K5"] = data["Attendance_percent"]

    ws["B9"] = data["total_late_days"]
    ws["E9"] = data["total_ot_hours"]
    ws["H9"] = data["avg_working_hours"]
    ws["K9"] = data["total_days"]


    plt.figure(figsize=(8,4))

    plt.plot(
        attendance_trend["Date"],
        attendance_trend["Present"],
        marker="o"
    )

    plt.title("Attendance Trend")
    plt.xlabel("Date")
    plt.ylabel("Present Employees")

    plt.xticks(rotation=45)

    plt.tight_layout()

    plt.savefig("Generated Reports/Charts/attendance_trend.png", dpi=300)

    plt.close()

    img = Image("Generated Reports/Charts/attendance_trend.png")

    img.width = 460
    img.height = 250
    ws.add_image(img, "A14")

    plt.figure(figsize=(7,4))

    plt.plot(
        monthly_ot["Month"].astype(str),
        monthly_ot["OT_Hours"],
        marker="o"
    )

    plt.title("Monthly Overtime Hours")
    plt.xlabel("Month")
    plt.ylabel("OT Hours")

    plt.xticks(rotation=45) 
    plt.grid(True)

    plt.tight_layout()

    plt.savefig("Generated Reports/Charts/monthly_ot.png")
    plt.close() 

    img = Image("Generated Reports/Charts/monthly_ot.png")

    img.width = 460
    img.height = 245
    ws.add_image(img, "H14")

    wb.save("Generated Reports/HR_Attendance_Report.xlsx")

df, employee_master = read_data()

df_long = clean_data(df,employee_master)

save_files(df_long)

data = calculate_kpis(df_long)

reports = calculate_summaries_n_exception(df_long)

monthly_ot, attendance_trend = generate_graphs(df_long)

generate_reports_n_Dashboard(
    reports,
    data,
    attendance_trend,
    monthly_ot
)

